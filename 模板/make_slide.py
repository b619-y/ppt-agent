#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
make_slide.py —— 按《风格规范.md》生成“结论标题 + 两条要点 + 左右两张论文图”的 16:9 幻灯片

用法：
    python3 make_slide.py slide_example.json            # 生成 pptx
    python3 make_slide.py slide_example.json --preview  # 另外用 LibreOffice 转 PDF，再用 pdftoppm 出 PNG 预览

配置（JSON）：单页直接写字段；多页写 {"output": "...", "slides": [ {单页字段}, ... ]}。
所有相对路径都相对于配置文件所在目录。要点里用 **双星号** 包住需要标红的片段。
版式参数（位置、字号、颜色）按《风格规范.md》中的当前规范定义，
一般不要改；确需微调时在配置里用 "style": {...} 覆盖同名键。
"""
import json
import os
import re
import shutil
import subprocess
import sys

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

STYLE = {
    "slide_w": 13.333, "slide_h": 7.5,
    "bullet_char": "❑",  # 符号使用 Segoe UI Symbol；普通文字的字体在 font() 中固定
    # 标题
    "title_box": [0.138, 0.177, 12.261, 0.673], "title_size": 40,
    # 红色横条（贯穿全宽）
    "bar_top": 0.98, "bar_h": 0.07,
    # 要点
    "bullets_box": [0.35, 1.18, 12.6, 1.05], "bullet_size": 24, "bullet_space_after": 4,
    # 图片区
    "img_top": 2.3, "img_margin": 0.25, "img_gap": 0.25, "img_max_h": 4.15,
    "img_border": "BFBFBF", "img_border_pt": 0.75,
    # 图注
    "caption_gap": 0.08, "caption_h": 0.8, "caption_size": 15, "source_size": 14,
    # 红框标注默认值
    "tag_size": 15, "tag_h": 0.36, "tag_line_pt": 2,
    # 页码
    "page_box": [12.6, 6.95, 0.5, 0.4], "page_size": 18,
    # 颜色（HEX）
    "red": "E60012", "blue": "1F3FE0", "black": "000000", "gray": "808080", "orange": "F4A460",
}


def rgb(hexstr):
    return RGBColor.from_string(hexstr.lstrip("#").upper())


class SlideBuilder:
    def __init__(self, prs, st):
        self.prs, self.st = prs, st
        self.s = prs.slides.add_slide(prs.slide_layouts[6])

    def font(self, run, size, color="000000", symbol=False):
        latin = "Segoe UI Symbol" if symbol else "Microsoft YaHei"
        f = run.font
        f.size, f.bold, f.name = Pt(size), False, latin
        f.color.rgb = rgb(color)
        rPr = run._r.get_or_add_rPr()
        for tag, name in (("a:ea", latin if symbol else "微软雅黑"), ("a:cs", latin)):
            el = rPr.find(qn(tag))
            if el is None:
                el = rPr.makeelement(qn(tag), {})
                rPr.append(el)
            el.set("typeface", name)

    def textbox(self, x, y, w, h, anchor=MSO_ANCHOR.TOP):
        tb = self.s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = False
        tf.margin_left = tf.margin_right = Inches(0.02)
        tf.margin_top = tf.margin_bottom = Inches(0)
        tf.vertical_anchor = anchor
        return tf

    # ---------- 各元素 ----------
    def title(self, text):
        st = self.st
        tf = self.textbox(*st["title_box"], anchor=MSO_ANCHOR.MIDDLE)
        r = tf.paragraphs[0].add_run(); r.text = text
        self.font(r, st["title_size"], st["black"])
        bar = self.s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(st["bar_top"]),
                                      self.prs.slide_width, Inches(st["bar_h"]))
        bar.fill.solid(); bar.fill.fore_color.rgb = rgb(st["red"])
        bar.line.fill.background(); bar.shadow.inherit = False

    def bullets(self, items):
        st = self.st
        tf = self.textbox(*st["bullets_box"])
        for i, line in enumerate(items):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.space_after = Pt(st["bullet_space_after"])
            r = p.add_run(); r.text = st["bullet_char"]
            self.font(r, st["bullet_size"], st["black"], symbol=True)
            # **...** 之间的片段标红
            for k, seg in enumerate(re.split(r"\*\*", line)):
                if not seg:
                    continue
                r = p.add_run(); r.text = seg
                self.font(r, st["bullet_size"], st["red"] if k % 2 else st["black"])

    def images(self, left, right):
        st = self.st
        wl, hl = Image.open(left["image"]).size
        wr, hr = Image.open(right["image"]).size
        al, ar = wl / hl, wr / hr
        avail = st["slide_w"] - 2 * st["img_margin"] - st["img_gap"]
        H = min(avail / (al + ar), st["img_max_h"])       # 两图等高，尽量大
        Wl, Wr = H * al, H * ar
        x0 = (st["slide_w"] - (Wl + Wr + st["img_gap"])) / 2
        top = st["img_top"]
        boxes = {"left": (x0, top, Wl, H), "right": (x0 + Wl + st["img_gap"], top, Wr, H)}
        for side, spec in (("left", left), ("right", right)):
            x, y, w, h = boxes[side]
            pic = self.s.shapes.add_picture(spec["image"], Inches(x), Inches(y), Inches(w), Inches(h))
            pic.line.color.rgb = rgb(st["img_border"]); pic.line.width = Pt(st["img_border_pt"])
            self.caption(x, w, spec, top + H)
            for t in spec.get("tags", []):
                self.tag(x, y, w, h, t)
            for m in spec.get("marks", []):
                self.mark(x, y, w, h, m)
        return H, Wl, Wr, x0

    def caption(self, x, w, spec, img_bottom):
        st = self.st
        pad, shift = spec.get("caption_pad", 0.0), spec.get("caption_shift", 0.0)
        tf = self.textbox(x - pad + shift, img_bottom + st["caption_gap"], w + 2 * pad, st["caption_h"])
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
        r = p.add_run(); r.text = spec["caption"]; self.font(r, st["caption_size"], st["blue"])
        p2 = tf.add_paragraph(); p2.alignment = PP_ALIGN.CENTER
        r = p2.add_run(); r.text = spec["source"]; self.font(r, st["source_size"], st["black"])

    def tag(self, x, y, w, h, t):
        """白底红框红字小标签。t: {"text", "x", "y"（占图宽/高的比例，左上角）, "w", "h"（英寸）, "size"}"""
        st = self.st
        b = self.s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x + w * t["x"]), Inches(y + h * t["y"]),
                                    Inches(t.get("w", 2.2)), Inches(t.get("h", st["tag_h"])))
        b.fill.solid(); b.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        b.line.color.rgb = rgb(st["red"]); b.line.width = Pt(st["tag_line_pt"]); b.shadow.inherit = False
        tf = b.text_frame; tf.word_wrap = False
        tf.margin_left = tf.margin_right = Inches(0.05); tf.margin_top = tf.margin_bottom = Inches(0)
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
        r = p.add_run(); r.text = t["text"]; self.font(r, t.get("size", st["tag_size"]), st["red"])

    def mark(self, x, y, w, h, m):
        """红色箭头/框选。坐标均为占图宽/高的比例。
        {"type": "box", "x", "y", "w", "h", "color": "orange"|"red"}
        {"type": "arrow", "x1", "y1", "x2", "y2", "double": true/false, "label": "GAP"}"""
        st = self.st
        color = st.get(m.get("color", "red"), m.get("color", st["red"]))
        if m["type"] == "box":
            b = self.s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x + w * m["x"]), Inches(y + h * m["y"]),
                                        Inches(w * m["w"]), Inches(h * m["h"]))
            b.fill.background(); b.line.color.rgb = rgb(color); b.line.width = Pt(m.get("pt", 2))
            b.shadow.inherit = False
        elif m["type"] == "arrow":
            c = self.s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
                                            Inches(x + w * m["x1"]), Inches(y + h * m["y1"]),
                                            Inches(x + w * m["x2"]), Inches(y + h * m["y2"]))
            c.line.color.rgb = rgb(color); c.line.width = Pt(m.get("pt", 2.5))
            ln = c.line._get_or_add_ln()
            if m.get("double"):
                ln.append(ln.makeelement(qn("a:headEnd"), {"type": "triangle", "w": "med", "len": "med"}))
            ln.append(ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"}))
            if m.get("label"):
                lx = x + w * (m["x1"] + m["x2"]) / 2 + m.get("label_dx", -0.9)
                ly = y + h * (m["y1"] + m["y2"]) / 2 - 0.25
                tf = self.textbox(lx, ly, 0.85, 0.5, MSO_ANCHOR.MIDDLE)
                p = tf.paragraphs[0]; p.alignment = PP_ALIGN.RIGHT
                r = p.add_run(); r.text = m["label"]; self.font(r, m.get("size", 20), color)

    def page_number(self, n):
        st = self.st
        tf = self.textbox(*st["page_box"], anchor=MSO_ANCHOR.MIDDLE)
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.RIGHT
        r = p.add_run(); r.text = str(n); self.font(r, st["page_size"], st["gray"])


def resolve(base, spec):
    spec = dict(spec)
    spec["image"] = os.path.join(base, spec["image"])
    if not os.path.exists(spec["image"]):
        sys.exit("找不到图片：%s" % spec["image"])
    return spec


def find_soffice():
    for c in ("soffice", "libreoffice", "/Applications/LibreOffice.app/Contents/MacOS/soffice"):
        p = shutil.which(c) or (c if os.path.exists(c) else None)
        if p:
            return p
    return None


def preview(pptx_path):
    so = find_soffice()
    if not so:
        print("未找到 LibreOffice：请直接用 PowerPoint/Keynote 打开 pptx 检查。")
        return
    outdir = os.path.dirname(pptx_path)
    subprocess.run([so, "--headless", "--convert-to", "pdf", "--outdir", outdir, pptx_path],
                   check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=180)
    pdf = os.path.splitext(pptx_path)[0] + ".pdf"
    stem = os.path.splitext(pptx_path)[0] + "_preview"
    if shutil.which("pdftoppm"):
        subprocess.run(["pdftoppm", "-r", "150", "-png", pdf, stem], check=True)
        # 单页时去掉 pdftoppm 自动加的页码后缀
        outs = sorted(f for f in os.listdir(outdir) if f.startswith(os.path.basename(stem)) and f.endswith(".png"))
        if len(outs) == 1 and outs[0] != os.path.basename(stem) + ".png":
            os.replace(os.path.join(outdir, outs[0]), stem + ".png")
        print("预览：%s*.png（请用眼睛逐页检查）" % stem)
    else:
        print("预览 PDF：%s（未安装 poppler，无法转 PNG）" % pdf)


def main():
    if len(sys.argv) < 2:
        sys.exit("用法：python3 make_slide.py <配置.json> [--preview]")
    cfg_path = sys.argv[1]
    base = os.path.dirname(os.path.abspath(cfg_path))
    cfg = json.load(open(cfg_path, encoding="utf-8"))
    st = dict(STYLE); st.update(cfg.get("style", {}))
    slides = cfg.get("slides", [cfg])

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(st["slide_w"]), Inches(st["slide_h"])
    for i, sc in enumerate(slides, 1):
        sst = dict(st); sst.update(sc.get("style", {}))
        b = SlideBuilder(prs, sst)
        b.title(sc["title"])
        b.bullets(sc["bullets"])
        H, Wl, Wr, x0 = b.images(resolve(base, sc["left"]), resolve(base, sc["right"]))
        b.page_number(sc.get("page", i))
        print("第 %d 页：图高 %.2f in，左图宽 %.2f in，右图宽 %.2f in，图注底边 %.2f in"
              % (i, H, Wl, Wr, sst["img_top"] + H + sst["caption_gap"] + sst["caption_h"]))
    out = os.path.join(base, cfg.get("output", "slide.pptx"))
    os.makedirs(os.path.dirname(out), exist_ok=True)
    prs.save(out)
    print("已生成：%s" % out)
    if "--preview" in sys.argv:
        preview(out)


if __name__ == "__main__":
    main()
